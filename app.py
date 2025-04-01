import os
import PyPDF2
import docx
import csv
import pandas as pd
import spacy
from fpdf import FPDF
from flask import Flask, request, jsonify, send_file, render_template
from werkzeug.utils import secure_filename

# Load NLP model
nlp = spacy.load('en_core_web_sm')

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads/'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Function to extract text from PDF
def extract_text_from_pdf(pdf_path):
    text = ""
    with open(pdf_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += page.extract_text() + "\n"
    return text

# Function to extract text from DOCX
def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    return "\n".join([para.text for para in doc.paragraphs])

# Function to extract key skills from text
def extract_skills(text):
    doc = nlp(text)
    skills = [token.text.lower() for token in doc if token.pos_ in ["NOUN", "PROPN", "ADJ"]]
    return set(skills)

# Function to calculate match score
def calculate_match_score(resume_text, job_desc):
    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_desc)

    match_score = (len(resume_skills & job_skills) / len(job_skills)) * 100 if job_skills else 0
    return round(match_score, 2)

# Generate structured PDF report
def generate_pdf(results):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", style='B', size=16)
    pdf.cell(200, 10, "Resume Analysis Report", ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", size=12)

    pdf.cell(100, 10, "Filename", border=1)
    pdf.cell(50, 10, "Score (%)", border=1)
    pdf.ln()

    for result in results:
        pdf.cell(100, 10, result['filename'], border=1)
        pdf.cell(50, 10, str(result['score']), border=1)
        pdf.ln()

    pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.pdf')
    pdf.output(pdf_path)
    return pdf_path

# Process single and batch resume uploads
@app.route('/', methods=['GET', 'POST'])
def upload_page():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_files():
    if 'files' not in request.files or 'job_desc' not in request.form:
        return jsonify({"error": "Missing files or job description"}), 400 
    
    files = request.files.getlist('files')
    job_desc = request.form['job_desc']
    results = []

    for file in files:
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)

        if filename.endswith('.pdf'):
            text = extract_text_from_pdf(file_path)
        elif filename.endswith('.docx'):
            text = extract_text_from_docx(file_path)
        else:
            continue

        score = calculate_match_score(text, job_desc)
        results.append({"filename": filename, "score": score})

    if not results:
        return jsonify({"error": "No valid resumes processed"}), 400 
    
    csv_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.csv')
    with open(csv_path, 'w', newline='') as csvfile:
        fieldnames = ['Filename', 'Score']
        writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
        writer.writeheader()
        for result in results:
            writer.writerow({'Filename': result['filename'], 'Score': result['score']})

    excel_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.xlsx')
    df = pd.DataFrame(results)
    df.to_excel(excel_path, index=False)

    pdf_path = generate_pdf(results)

    return jsonify({
        "results": results,
        "csv_download": "/download_csv",
        "excel_download": "/download_excel",
        "pdf_download": "/download_pdf"
    })

@app.route('/download_csv', methods=['GET'])
def download_csv():
    csv_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.csv')
    return send_file(csv_path, as_attachment=True)

@app.route('/download_excel', methods=['GET'])
def download_excel():
    excel_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.xlsx')
    return send_file(excel_path, as_attachment=True)

@app.route('/download_pdf', methods=['GET'])
def download_pdf():
    pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], 'resume_results.pdf')
    return send_file(pdf_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
